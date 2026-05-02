from pathlib import Path

import docx
import openpyxl
from pypdf import PdfReader


def extract_text(file_path: str) -> str:
    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix in {".txt", ".md", ".csv"}:
        return path.read_text(encoding="utf-8", errors="ignore")

    if suffix == ".docx":
        document = docx.Document(str(path))
        return "\n".join(paragraph.text for paragraph in document.paragraphs)

    if suffix == ".xlsx":
        workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        lines = []
        for sheet in workbook.worksheets:
            lines.append(f"# Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = [str(cell) for cell in row if cell is not None]
                if values:
                    lines.append(" | ".join(values))
        return "\n".join(lines)

    if suffix == ".pdf":
        reader = PdfReader(str(path))
        return "\n".join((page.extract_text() or "") for page in reader.pages)

    raise ValueError(f"Formato não suportado: {suffix}")
